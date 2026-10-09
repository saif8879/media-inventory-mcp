import io
import json
from pathlib import Path
from unittest.mock import AsyncMock, MagicMock, patch

from botocore.exceptions import ClientError
from httpx import ASGITransport, AsyncClient
from PIL import Image
import pptx
from pptx import Presentation
import pytest

from mcp_server.server import create_app, create_mcp_server
from mcp_server.services.deck_service import DeckBuilderService, format_inr


def test_deck_dependencies_and_exports_dir():
    assert pptx.__version__ is not None
    assert Image.__version__ is not None


def test_compute_contain_fit():
    service = DeckBuilderService()
    # Wide image: 1600x800 into 8.0 x 6.0 box
    x, y, w, h = service.compute_contain_fit(1600, 800, box_x=1.0, box_y=1.0, box_w=8.0, box_h=6.0)
    assert round(w, 2) == 8.0
    assert round(h, 2) == 4.0
    assert round(x, 2) == 1.0
    assert round(y, 2) == 2.0  # Centered vertically in 6.0 box

    # Tall image: 800x1600 into 8.0 x 6.0 box
    x2, y2, w2, h2 = service.compute_contain_fit(800, 1600, box_x=1.0, box_y=1.0, box_w=8.0, box_h=6.0)
    assert round(h2, 2) == 6.0
    assert round(w2, 2) == 3.0
    assert round(x2, 2) == 3.5  # Centered horizontally in 8.0 box

    # Edge cases: 0 or negative dimensions should fall back to box dimensions
    x3, y3, w3, h3 = service.compute_contain_fit(0, 0, box_x=1.0, box_y=1.0, box_w=8.0, box_h=6.0)
    assert (x3, y3, w3, h3) == (1.0, 1.0, 8.0, 6.0)


def test_build_deck_creates_valid_pptx():
    service = DeckBuilderService()
    sample_sites = [
        {
            "id": "site-1",
            "inventory_code": "MI1",
            "city": "Mumbai",
            "area": "Bandra West",
            "location": "Linking Road, Bandra",
            "media_type": "Billboard",
            "sub_type": "Frontlit",
            "facing": "North",
            "width": 30.0,
            "height": 15.0,
            "area_sq_ft": 450.0,
            "total_cost": 450000.0,
            "image_url": None,
            "admavin": {
                "reach": 1250000,
                "ots": 3400000,
                "lts_score": 0.88,
            },
        }
    ]
    deck_bytes, summary = service.build_deck(
        sites=sample_sites,
        campaign_title="Q4 Festival Campaign",
        client_name="Test Brand",
        theme="editorial",
        include_commercials=True,
        include_audience_metrics=True,
    )
    assert len(deck_bytes) > 0
    assert summary["site_count"] == 1
    assert summary["total_monthly_cost"] == 450000.0
    assert summary["total_reach"] == 1250000
    assert summary["cities"] == ["Mumbai"]

    # Verify PPTX is readable
    prs = Presentation(io.BytesIO(deck_bytes))
    # 3 slides: Cover, Site 1, Closing
    assert len(prs.slides) == 3
    # Check 16:9 widescreen dimensions (13.333" x 7.50")
    assert round(prs.slide_width.inches, 3) == 13.333
    assert round(prs.slide_height.inches, 2) == 7.50


def test_build_deck_classic_theme_and_multiple_sites():
    service = DeckBuilderService()
    sample_sites = [
        {
            "id": "site-1",
            "inventory_code": "MI1",
            "city": "Mumbai",
            "area": "Bandra West",
            "location": "Linking Road",
            "media_type": "Billboard",
            "sub_type": "Frontlit",
            "facing": "North",
            "width": 30.0,
            "height": 15.0,
            "area_sq_ft": 450.0,
            "total_cost": 250000.0,
            "image_url": None,
        },
        {
            "id": "site-2",
            "inventory_code": "MI2",
            "city": "Delhi",
            "area": "Connaught Place",
            "location": "Inner Circle",
            "media_type": "Unipole",
            "sub_type": "Backlit",
            "facing": "South",
            "width": 40.0,
            "height": 20.0,
            "area_sq_ft": 800.0,
            "total_cost": 350000.0,
            "image_url": None,
        },
    ]
    deck_bytes, summary = service.build_deck(
        sites=sample_sites,
        campaign_title="National Brand Rollout",
        client_name="Acme Corp",
        theme="classic",
        include_commercials=False,
        include_audience_metrics=False,
    )
    assert len(deck_bytes) > 0
    assert summary["site_count"] == 2
    assert summary["total_monthly_cost"] == 600000.0
    assert sorted(summary["cities"]) == ["Delhi", "Mumbai"]

    prs = Presentation(io.BytesIO(deck_bytes))
    # 4 slides: Cover, Site 1, Site 2, Closing
    assert len(prs.slides) == 4


def test_build_deck_with_embedded_image():
    service = DeckBuilderService()
    # Create a small valid test image in-memory
    img = Image.new("RGB", (640, 480), color=(100, 150, 200))
    buf = io.BytesIO()
    img.save(buf, format="JPEG")
    img_bytes = buf.getvalue()

    sample_sites = [
        {
            "id": "site-img",
            "inventory_code": "MI-IMG",
            "city": "Bengaluru",
            "area": "Indiranagar",
            "location": "100ft Road",
            "media_type": "Billboard",
            "sub_type": "Digital",
            "facing": "East",
            "width": 20.0,
            "height": 10.0,
            "area_sq_ft": 200.0,
            "total_cost": 150000.0,
            "image_url": "https://example.com/photo.jpg",
        }
    ]

    with patch.object(service, "fetch_image", return_value=(io.BytesIO(img_bytes), (640, 480))):
        deck_bytes, summary = service.build_deck(
            sites=sample_sites,
            campaign_title="Bengaluru Tech Tour",
            client_name="Tech Co",
            theme="editorial",
        )
        assert len(deck_bytes) > 0
        prs = Presentation(io.BytesIO(deck_bytes))
        assert len(prs.slides) == 3
        # Site slide is slide index 1
        site_slide = prs.slides[1]
        # Verify at least one picture shape was added
        picture_shapes = [s for s in site_slide.shapes if s.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.PICTURE]
        assert len(picture_shapes) >= 1


def test_save_and_upload_success(tmp_path):
    mock_s3 = MagicMock()
    mock_s3.put_object.return_value = {}
    mock_s3.generate_presigned_url.return_value = "https://s3.ap-south-1.amazonaws.com/test-bucket/campaign-decks/campaign_deck_test.pptx"

    service = DeckBuilderService(s3_client=mock_s3, bucket="test-bucket", exports_dir=tmp_path)
    result = service.save_and_upload(
        deck_bytes=b"fake-pptx-data",
        file_slug="mumbai_launch",
        expires_in=600,
    )

    assert result["file_name"].startswith("campaign_deck_mumbai_launch_")
    assert result["file_name"].endswith(".pptx")
    assert Path(result["local_file_path"]).exists()
    assert Path(result["local_file_path"]).read_bytes() == b"fake-pptx-data"
    assert result["s3_download_url"] == "https://s3.ap-south-1.amazonaws.com/test-bucket/campaign-decks/campaign_deck_test.pptx"
    assert result["expires_in_seconds"] == 600

    mock_s3.put_object.assert_called_once()
    mock_s3.generate_presigned_url.assert_called_once()


def test_save_and_upload_s3_error_handled(tmp_path):
    mock_s3 = MagicMock()
    mock_s3.put_object.side_effect = ClientError({"Error": {"Code": "AccessDenied", "Message": "Access Denied"}}, "PutObject")

    service = DeckBuilderService(s3_client=mock_s3, bucket="test-bucket", exports_dir=tmp_path)
    result = service.save_and_upload(
        deck_bytes=b"fake-pptx-data",
        file_slug="error_test",
        expires_in=600,
    )

    assert Path(result["local_file_path"]).exists()
    assert result["s3_download_url"] is None


def test_format_inr_numbers():
    assert format_inr(None) == "—"
    assert format_inr(0) == "0"
    assert format_inr(500) == "500"
    assert format_inr(450000) == "4,50,000"
    assert format_inr(10000000) == "1,00,00,000"
    assert format_inr(-500) == "-500"
    assert format_inr(-450000) == "-4,50,000"


def test_fetch_image_bucket_none():
    service = DeckBuilderService()
    service.bucket = None

    # 1. With HTTP URL that fails to download or returns 404 (mock httpx)
    with patch("httpx.Client") as mock_client_cls:
        mock_client = MagicMock()
        mock_client.__enter__.return_value = mock_client
        mock_resp = MagicMock()
        mock_resp.status_code = 404
        mock_client.get.return_value = mock_resp
        mock_client_cls.return_value = mock_client

        img_stream, dims = service.fetch_image("https://example.com/test.jpg")
        assert img_stream is None
        assert dims is None

    # 2. With S3 key string when bucket is None
    img_stream_key, dims_key = service.fetch_image("inventory/site1.jpg")
    assert img_stream_key is None
    assert dims_key is None


def test_fetch_image_s3_url_unquotes_key():
    mock_s3 = MagicMock()
    img = Image.new("RGB", (100, 100), color=(10, 20, 30))
    buf = io.BytesIO()
    img.save(buf, format="JPEG")
    img_bytes = buf.getvalue()

    mock_body = MagicMock()
    mock_body.read.return_value = img_bytes
    mock_s3.get_object.return_value = {"Body": mock_body}

    service = DeckBuilderService(s3_client=mock_s3, bucket="my-bucket")
    url = "https://my-bucket.s3.ap-south-1.amazonaws.com/my-bucket/site%20photos/photo%201.jpg"
    stream, dims = service.fetch_image(url)

    assert stream is not None
    assert dims == (100, 100)
    mock_s3.get_object.assert_called_once_with(
        Bucket="my-bucket",
        Key="site photos/photo 1.jpg",
    )


@pytest.mark.asyncio
async def test_mcp_server_lists_generate_campaign_deck():
    server = create_mcp_server()
    tools = await server.list_tools()
    tool_names = {t.name for t in tools}
    assert "generate_campaign_deck" in tool_names


@pytest.mark.asyncio
async def test_generate_campaign_deck_handler_no_sites():
    server = create_mcp_server()
    res = await server.call_tool(
        "generate_campaign_deck",
        {"site_identifiers": []}
    )
    assert not res.is_error
    data = json.loads(res.content[0].text)
    assert data["success"] is False
    assert data["error"]["code"] == "NO_VALID_SITES"


@pytest.mark.asyncio
async def test_generate_campaign_deck_handler_all_invalid_sites():
    server = create_mcp_server()
    with patch("mcp_server.tools.generate_campaign_deck.MediaInventoryService") as mock_srv_cls:
        mock_srv = mock_srv_cls.return_value
        mock_srv.get_media_by_identifier = AsyncMock(return_value=None)
        res = await server.call_tool(
            "generate_campaign_deck",
            {"site_identifiers": ["INVALID_CODE_1", "INVALID_CODE_2"]}
        )
        assert not res.is_error
        data = json.loads(res.content[0].text)
        assert data["success"] is False
        assert data["error"]["code"] == "NO_VALID_SITES"
        assert len(data["error"]["details"]["warnings"]) == 2


@pytest.mark.asyncio
async def test_generate_campaign_deck_handler_success_with_warnings(tmp_path):
    server = create_mcp_server()
    site_valid = {
        "id": "site-1",
        "inventory_code": "MI1",
        "city": "Mumbai",
        "area": "Bandra West",
        "location": "Linking Road",
        "media_type": "Billboard",
        "sub_type": "Frontlit",
        "facing": "North",
        "width": 30.0,
        "height": 15.0,
        "area_sq_ft": 450.0,
        "total_cost": 450000.0,
        "image_url": None,
        "admavin": {
            "reach": 1250000,
            "ots": 3400000,
            "lts_score": 0.88,
        },
    }

    async def mock_get_by_id(identifier, **kwargs):
        if identifier == "MI1":
            return site_valid
        return None

    with patch("mcp_server.tools.generate_campaign_deck.MediaInventoryService") as mock_srv_cls, \
         patch("mcp_server.tools.generate_campaign_deck.DeckBuilderService") as mock_deck_cls:
        mock_srv = mock_srv_cls.return_value
        mock_srv.get_media_by_identifier = AsyncMock(side_effect=mock_get_by_id)

        mock_deck_instance = MagicMock()
        mock_deck_cls.return_value = mock_deck_instance
        mock_deck_instance.build_deck.return_value = (b"fake-pptx", {
            "site_count": 1,
            "total_monthly_cost": 450000.0,
            "total_reach": 1250000,
            "cities": ["Mumbai"],
        })
        mock_deck_instance.save_and_upload.return_value = {
            "file_name": "campaign_deck_proposal_test.pptx",
            "local_file_path": str(tmp_path / "campaign_deck_proposal_test.pptx"),
            "s3_download_url": "https://s3.example.com/test.pptx",
            "expires_in_seconds": 900,
        }

        res = await server.call_tool(
            "generate_campaign_deck",
            {
                "site_identifiers": ["MI1", "NON_EXISTENT"],
                "campaign_title": "Summer Festival",
                "client_name": "Pepsi",
                "theme": "editorial",
                "include_commercials": True,
                "include_audience_metrics": True,
                "expires_in": 900,
            }
        )
        assert not res.is_error
        data = json.loads(res.content[0].text)
        assert data["success"] is True
        deck_data = data["data"]
        assert deck_data["site_count"] == 1
        assert deck_data["theme"] == "editorial"
        assert deck_data["file_name"] == "campaign_deck_proposal_test.pptx"
        assert deck_data["s3_download_url"] == "https://s3.example.com/test.pptx"
        assert len(deck_data["warnings"]) == 1
        assert "NON_EXISTENT" in deck_data["warnings"][0]
        mock_deck_instance.build_deck.assert_called_once_with(
            sites=[site_valid],
            campaign_title="Summer Festival",
            client_name="Pepsi",
            theme="editorial",
            include_commercials=True,
            include_audience_metrics=True,
        )
        mock_deck_instance.save_and_upload.assert_called_once_with(
            deck_bytes=b"fake-pptx",
            file_slug="Summer Festival",
            expires_in=900,
        )


@pytest.mark.asyncio
async def test_health_endpoint_lists_generate_campaign_deck():
    app = create_app()
    async with AsyncClient(
        transport=ASGITransport(app=app), base_url="http://test"
    ) as client:
        res = await client.get("/health")
        assert res.status_code == 200
        health_data = res.json()
        assert "generate_campaign_deck" in health_data["tools"]


